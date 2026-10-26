"""
Resume reader + layout audit.

Produces two texts from the file:
  * `text`: the human reading order (column-aware), used for profile and
    keyword extraction.
  * `naive_text`: what a column-unaware parser produces by reading straight
    across the page, interleaving the columns line by line.
and a layout audit of the structural features that ATS parsers commonly
mishandle, measured from the document itself rather than guessed.
"""
import io
import re
from typing import Any, Dict, List, Optional, Tuple

import pdfplumber
import docx

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?91[\s-]?)?(?<!\d)[6-9]\d{4}[\s-]?\d{5}(?!\d)|\+\d{1,3}[\s-]?\(?\d{2,4}\)?[\s-]?\d{3,4}[\s-]?\d{3,4}")
LINKEDIN_RE = re.compile(r"linkedin\.com/in/[\w\-%]+", re.IGNORECASE)
GITHUB_RE = re.compile(r"github\.com/[\w\-]+", re.IGNORECASE)
PUA_RE = re.compile(r"[-]")
EMOJI_RE = re.compile(r"[☀-➿\U0001F300-\U0001FAFF☎✉■-◿]")

STANDARD_SECTIONS = {
    "summary": ["summary", "professional summary", "profile", "career objective", "objective", "about", "career summary", "profile summary"],
    "experience": ["experience", "work experience", "professional experience", "employment history", "work history", "employment",
                   "internships", "internship", "internship experience", "relevant experience", "industry experience"],
    "education": ["education", "academic background", "academics", "educational qualifications", "academic qualifications",
                  "qualifications", "education details", "academic details", "scholastic achievements"],
    "skills": ["skills", "technical skills", "core competencies", "key skills", "skills summary", "technologies", "technical proficiency",
               "skills & tools", "skills and tools", "tools", "competencies"],
    "projects": ["projects", "academic projects", "personal projects", "key projects", "project experience", "major projects"],
    "certifications": ["certifications", "certificates", "licenses & certifications", "courses", "certifications & courses",
                       "training", "trainings"],
    "achievements": ["achievements", "awards", "honors", "honours", "accomplishments", "awards & achievements", "awards and achievements"],
    "leadership": ["positions of responsibility", "leadership", "leadership experience", "extracurricular activities",
                   "extra-curricular activities", "co-curricular activities", "volunteering", "volunteer experience", "activities"],
    "publications": ["publications", "research", "research experience", "papers"],
    "other": ["languages", "interests", "hobbies", "personal details", "personal information", "declaration", "references"],
}
NONSTANDARD_SECTIONS = {
    "what i've built": "projects", "things i've built": "projects", "my work": "projects", "portfolio": "projects", "builds": "projects",
    "my journey": "experience", "where i've worked": "experience", "career journey": "experience", "work": "experience",
    "superpowers": "skills", "tech stack": "skills", "toolbox": "skills", "toolkit": "skills", "what i know": "skills", "expertise": "skills",
    "about me": "summary", "who i am": "summary", "hello": "summary", "background": "summary",
    "learning": "education", "where i studied": "education",
}
_HEADER_LOOKUP = {h: canon for canon, hs in STANDARD_SECTIONS.items() for h in hs}

PERSONAL_DETAIL_PATTERNS = {
    "Date of birth": r"\b(?:date of birth|d\.?o\.?b\.?)\b",
    "Gender": r"\bgender\s*[:\-]",
    "Marital status": r"\bmarital status\b",
    "Religion": r"\breligion\s*[:\-]",
    "Father's name": r"\bfather'?s name\b",
    "Nationality": r"\bnationality\s*[:\-]",
}


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

X_TOL = 1.5


def _group_lines(words: List[Dict[str, Any]], tol: float = 3.0) -> List[List[Dict[str, Any]]]:
    lines: List[List[Dict[str, Any]]] = []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if lines and abs(lines[-1][0]["top"] - w["top"]) <= tol:
            lines[-1].append(w)
        else:
            lines.append([w])
    return [sorted(l, key=lambda w: w["x0"]) for l in lines]


def _line_text(line: List[Dict[str, Any]]) -> str:
    out, prev = "", None
    for w in line:
        if prev is not None:
            out += "   " if w["x0"] - prev["x1"] > 25 else " "
        out += w["text"]
        prev = w
    return out


def _find_gutter(words: List[Dict[str, Any]], width: float) -> Optional[float]:
    """x-position of an empty vertical channel that separates two text columns."""
    if len(words) < 60:
        return None
    lines = _group_lines(words)
    if len(lines) < 12:
        return None
    bin_w = 2.0
    nbins = int(width / bin_w) + 1
    cover = [0] * nbins
    for line in lines:
        hit = set()
        for w in line:
            for b in range(int(w["x0"] / bin_w), min(nbins, int(w["x1"] / bin_w) + 1)):
                hit.add(b)
        for b in hit:
            cover[b] += 1
    lo, hi = int(0.22 * nbins), int(0.78 * nbins)
    limit = max(2, int(0.12 * len(lines)))  # a name/title banner may span both columns
    best: Tuple[int, int] = (0, 0)
    run_start = None
    for b in range(lo, hi + 1):
        if b < hi and cover[b] <= limit:
            if run_start is None:
                run_start = b
        else:
            if run_start is not None and b - run_start > best[1] - best[0]:
                best = (run_start, b)
            run_start = None
    if (best[1] - best[0]) * bin_w < 8:
        return None
    gutter = (best[0] + best[1]) / 2 * bin_w
    left = sum(1 for w in words if w["x1"] <= gutter)
    right = sum(1 for w in words if w["x0"] >= gutter)
    if left < 0.15 * len(words) or right < 0.15 * len(words):
        return None
    # Right-aligned dates in a single-column resume also leave a gap, but only
    # on a few lines; genuine columns put text on both sides of most lines.
    with_left = sum(1 for l in lines if any(w["x1"] <= gutter for w in l))
    with_right = sum(1 for l in lines if any(w["x0"] >= gutter for w in l))
    if min(with_left, with_right) < 0.3 * len(lines):
        return None
    return gutter


def _column_text(words: List[Dict[str, Any]], gutter: float) -> str:
    lines = _group_lines(words)
    span, left, right = [], [], []
    for line in lines:
        l = [w for w in line if w["x1"] <= gutter]
        r = [w for w in line if w["x0"] >= gutter]
        if len(l) + len(r) < len(line):
            span.append(line)
        else:
            if l:
                left.append(l)
            if r:
                right.append(r)
    col_top = min([l[0]["top"] for l in left + right] or [0])
    top_span = [s for s in span if s[0]["top"] < col_top]
    rest_span = [s for s in span if s[0]["top"] >= col_top]
    ordered = top_span + left + right + rest_span
    return "\n".join(_line_text(l) for l in ordered)


def _read_pdf(data: bytes) -> Dict[str, Any]:
    reading, naive = [], []
    layout = {"pages": 0, "column_pages": 0, "tables": 0, "images": 0, "icons": 0, "photo": False,
              "margin_contact": [], "fonts": set(), "cid_garbage": False}
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        layout["pages"] = len(pdf.pages)
        for page in pdf.pages:
            # LaTeX and tightly kerned PDFs put ~2pt between words and emit no
            # space characters; pdfplumber's default 3pt tolerance glues them
            # into "ProductManagement", which breaks every keyword match.
            words = page.extract_words(x_tolerance=X_TOL, keep_blank_chars=False, use_text_flow=False, extra_attrs=["fontname"])
            for w in words[:400]:
                layout["fonts"].add(re.sub(r"^[A-Z]{6}\+", "", w.get("fontname", "")).split("-")[0])
            naive.append(page.extract_text(x_tolerance=X_TOL, layout=False) or "")
            gutter = _find_gutter(words, float(page.width))
            if gutter:
                layout["column_pages"] += 1
                reading.append(_column_text(words, gutter))
            else:
                reading.append("\n".join(_line_text(l) for l in _group_lines(words)))
            try:
                tables = [t for t in page.find_tables() if len(t.rows) >= 2 and len(t.rows[0].cells) >= 2]
                layout["tables"] += len(tables)
            except Exception:
                pass
            for img in page.images:
                w, h = float(img["x1"] - img["x0"]), float(img["bottom"] - img["top"])
                if w <= 26 and h <= 26:
                    layout["icons"] += 1
                else:
                    layout["images"] += 1
                    if w >= 50 and h >= 50 and img["top"] < page.height * 0.3:
                        layout["photo"] = True
            h = float(page.height)
            for line in _group_lines(words):
                if line[0]["top"] < 30 or line[0]["bottom"] > h - 30:
                    t = _line_text(line)
                    if EMAIL_RE.search(t) or PHONE_RE.search(t):
                        layout["margin_contact"].append(t.strip())
    text = "\n".join(reading)
    layout["cid_garbage"] = "(cid:" in text or "(cid:" in "".join(naive)
    layout["fonts"] = sorted(f for f in layout["fonts"] if f)
    return {"text": text, "naive_text": "\n".join(naive), "layout": layout, "file_type": "pdf"}


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def _read_docx(data: bytes) -> Dict[str, Any]:
    d = docx.Document(io.BytesIO(data))
    body = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = []
            for c in row.cells:
                if c.text.strip() and c.text.strip() not in cells:
                    cells.append(c.text.strip())
            if cells:
                body.append("   ".join(cells))
    header_text, footer_text = [], []
    for s in d.sections:
        header_text += [p.text for p in s.header.paragraphs if p.text.strip()]
        footer_text += [p.text for p in s.footer.paragraphs if p.text.strip()]
    xml = d.element.xml
    textbox_texts = re.findall(r"<w:txbxContent.*?</w:txbxContent>", xml, re.DOTALL)
    textbox_lines = []
    for tb in textbox_texts:
        t = " ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", tb)).strip()
        if t:
            textbox_lines.append(t)
    cols = re.findall(r'<w:cols[^>]*w:num="(\d+)"', xml)
    layout = {
        "pages": None, "column_pages": 1 if any(int(c) > 1 for c in cols) else 0,
        "tables": len(d.tables), "images": len(d.inline_shapes), "icons": 0, "photo": False,
        "margin_contact": [t for t in header_text + footer_text if EMAIL_RE.search(t) or PHONE_RE.search(t)],
        "header_text": header_text, "footer_text": footer_text, "text_boxes": len(textbox_lines),
        "fonts": [], "cid_garbage": False,
    }
    full = "\n".join(header_text + body + textbox_lines + footer_text)
    return {"text": full, "naive_text": "\n".join(body), "layout": layout, "file_type": "docx"}


# --------------------------------------------------------------------------
# Audit
# --------------------------------------------------------------------------

def detect_sections(text: str) -> Dict[str, Any]:
    found: Dict[str, str] = {}
    nonstandard: List[Dict[str, str]] = []
    for raw in text.split("\n"):
        line = raw.strip().strip(":").strip()
        if not line or len(line) > 42 or len(line.split()) > 5:
            continue
        key = re.sub(r"\s+", " ", re.sub(r"[^a-z&' ]", " ", line.lower())).strip()
        if key in _HEADER_LOOKUP:
            found.setdefault(_HEADER_LOOKUP[key], line)
        elif key in NONSTANDARD_SECTIONS and (line.isupper() or line.istitle() or len(line.split()) <= 3):
            nonstandard.append({"header": line, "suggested": NONSTANDARD_SECTIONS[key].title()})
    return {"found": found, "nonstandard": nonstandard}


def audit(read: Dict[str, Any]) -> Dict[str, Any]:
    text = read["text"]
    lay = read["layout"]
    words = len(re.findall(r"\w+", text))
    sections = detect_sections(text)
    contact = {
        "email": (EMAIL_RE.search(text) or [None])[0] if EMAIL_RE.search(text) else None,
        "phone": PHONE_RE.search(text).group(0).strip() if PHONE_RE.search(text) else None,
        "linkedin": LINKEDIN_RE.search(text).group(0) if LINKEDIN_RE.search(text) else None,
        "github": GITHUB_RE.search(text).group(0) if GITHUB_RE.search(text) else None,
    }
    personal = [label for label, pat in PERSONAL_DETAIL_PATTERNS.items() if re.search(pat, text, re.IGNORECASE)]
    bullet_lines = [l for l in text.split("\n") if re.match(r"\s*[•\-\*▪●◦–]\s+\S", l)]
    quantified = [l for l in bullet_lines if re.search(r"\d+(?:\.\d+)?\s*(?:%|x\b|[kmb]\b|mn\b|cr\b|\+)|₹|\$|\b\d{2,}\b|\blakh|\bcrore", l, re.IGNORECASE)]

    issues: List[Dict[str, Any]] = []

    def add(key, severity, title, detail, fix):
        issues.append({"key": key, "severity": severity, "title": title, "detail": detail, "fix": fix})

    scanned = read["file_type"] == "pdf" and words < 60
    if scanned:
        add("scanned", "critical", "Resume is an image, not text",
            "Almost no selectable text was found, so an ATS sees a blank document.",
            "Export the resume from Word or Google Docs as a text PDF. Don't scan or screenshot it.")
    if lay.get("cid_garbage"):
        add("encoding", "critical", "Broken font encoding",
            "Parts of the text come out as '(cid:NN)' codes, so those words are unreadable to a parser.",
            "Re-export with a standard font (Calibri, Arial, Garamond), or use 'Save as PDF' instead of 'Print to PDF'.")
    if lay.get("column_pages"):
        add("columns", "high", "Multi-column layout",
            f"{lay['column_pages']} page(s) use side-by-side columns. Parsers that read straight across the page will interleave them.",
            "Move to a single-column layout. Put skills in their own section under the summary, not in a sidebar.")
    if lay.get("tables"):
        add("tables", "medium", f"{lay['tables']} table(s) detected",
            "Many parsers flatten tables into run-on text or skip cells entirely.",
            "Replace tables with plain lines, e.g. 'B.Tech, Computer Science, XYZ University, 2024, CGPA 8.6'.")
    if lay.get("text_boxes"):
        add("text_boxes", "high", f"{lay['text_boxes']} text box(es)",
            "Text inside Word text boxes is often skipped by parsers.",
            "Move that content into the normal document body.")
    pua = len(PUA_RE.findall(text))
    if lay.get("icons") or pua:
        add("icons", "low", "Icons used as labels",
            f"{lay.get('icons', 0) + pua} icon glyph(s) found. Parsers drop icons, so any label that is only an icon disappears.",
            "Write the label as text ('Email:', 'Phone:') or leave it out.")
    if lay.get("margin_contact"):
        add("header_contact", "high", "Contact details in the page header/footer",
            "Several enterprise ATS (Workday, Taleo, iCIMS) skip header/footer areas, so your email or phone can be lost.",
            "Put your name, email and phone as normal text in the first lines of the body.")
    if not contact["email"]:
        add("no_email", "high", "No email address found", "Recruiters can't contact you, and many ATS require an email to create your profile.",
            "Add a professional email in the first lines.")
    if not contact["phone"]:
        add("no_phone", "medium", "No phone number found", "Indian recruiters usually call first.", "Add a phone number with the +91 country code.")
    renamed = {ns["suggested"].lower() for ns in sections["nonstandard"]}
    for essential in ("experience", "education", "skills"):
        if essential not in sections["found"] and essential not in renamed and not scanned:
            if essential == "experience" and "projects" in sections["found"]:
                continue
            add(f"missing_{essential}", "medium", f"No '{essential.title()}' heading found",
                f"ATS map content into fields by section heading; without a clear '{essential.title()}' heading this content may be misfiled.",
                f"Add a plain heading: '{essential.title()}'.")
    for ns in sections["nonstandard"]:
        add("nonstandard_header", "medium", f"Creative heading '{ns['header']}'",
            "Strict parsers don't recognise creative section names, and the content under them can be misfiled.",
            f"Rename it to '{ns['suggested']}'.")
    if personal or lay.get("photo"):
        items = personal + (["Photo"] if lay.get("photo") else [])
        add("personal_details", "medium", "Personal details that invite bias",
            f"Found: {', '.join(items)}. ATS don't use these, and they can enable bias in manual screening.",
            "Remove date of birth, gender, marital status, religion, father's name and photo.")
    if lay.get("pages") and lay["pages"] > 2:
        add("length", "low", f"{lay['pages']} pages long", "Recruiters skim. For under 8 years of experience, 1–2 pages works best.", "Trim it to 1–2 pages.")
    if 0 < words < 180 and not scanned:
        add("thin", "medium", "Very little content", f"Only about {words} words, which leaves few keywords for an ATS to match.",
            "Add project and experience bullets that name the tools you used.")
    if len(bullet_lines) >= 4 and len(quantified) / len(bullet_lines) < 0.3:
        add("quantify", "low", "Few measurable results",
            f"Only {len(quantified)} of {len(bullet_lines)} bullets include a number.",
            "Add impact numbers: '% faster', 'users', '₹ saved', 'requests/sec'.")
    if len(lay.get("fonts") or []) > 4:
        add("fonts", "low", f"{len(lay['fonts'])} different fonts", "Mixed fonts can come from copy-pasting and sometimes carry encoding problems.",
            "Use one or two fonts throughout.")

    return {
        "file_type": read["file_type"], "pages": lay.get("pages"), "words": words,
        "columns": bool(lay.get("column_pages")), "tables": lay.get("tables", 0), "images": lay.get("images", 0),
        "icons": lay.get("icons", 0) + pua, "text_boxes": lay.get("text_boxes", 0), "photo": lay.get("photo", False),
        "margin_contact": lay.get("margin_contact", []), "cid_garbage": lay.get("cid_garbage", False), "scanned": scanned,
        "fonts": lay.get("fonts", []), "sections": sections["found"], "nonstandard_sections": sections["nonstandard"],
        "contact": contact, "personal_details": personal, "bullets": len(bullet_lines), "quantified_bullets": len(quantified),
        "issues": issues,
    }


_PENALTIES = {
    "scanned": (70, 70), "encoding": (35, 35), "text_boxes": (18, 18), "no_email": (10, 10), "no_phone": (4, 4),
    "missing_experience": (10, 6), "missing_education": (8, 5), "missing_skills": (8, 5), "thin": (6, 6),
    "personal_details": (0, 0), "length": (0, 0), "quantify": (0, 0), "fonts": (2, 2),
}


def parse_score(audit_result: Dict[str, Any], profile: Dict[str, Any]) -> Dict[str, Any]:
    """How much of the resume survives this ATS's parser (0-100), with the reasons."""
    p = profile["parsing"]
    score = 100.0
    reasons = []
    for issue in audit_result["issues"]:
        k = issue["key"]
        if k == "columns":
            pen = 24 if not p["handles_columns"] else 5
        elif k == "tables":
            pen = min(16, 6 * audit_result["tables"]) if not p["handles_tables"] else 3
        elif k == "icons":
            pen = 5 if p["drops_icons"] else 1
        elif k == "header_contact":
            pen = 16 if p["drops_header_footer"] else 2
        elif k == "nonstandard_header":
            pen = 7 if p["strict_headers"] else 2
        elif k in _PENALTIES:
            strict, lenient = _PENALTIES[k]
            pen = strict if p["strict_headers"] else lenient
        else:
            pen = 0
        if pen:
            score -= pen
            reasons.append({"issue": issue["title"], "penalty": pen})
    return {"score": int(max(5, min(100, round(score)))), "reasons": reasons}


def ats_view(read: Dict[str, Any], audit_result: Dict[str, Any], profile: Dict[str, Any]) -> Dict[str, Any]:
    """Simulates the text this ATS's parser would extract, with every change annotated."""
    p = profile["parsing"]
    base = read["naive_text"] if (audit_result["columns"] and not p["handles_columns"]) else read["text"]
    changes = []
    if audit_result["columns"] and not p["handles_columns"]:
        changes.append({"type": "columns", "text": "Left and right columns read across as single lines",
                        "reason": f"{profile['name']} reads multi-column pages line by line across the full width."})
    margin = set(audit_result.get("margin_contact") or [])
    lines = []
    for line in base.split("\n"):
        s = line.strip()
        if not s:
            continue
        if p["drops_header_footer"] and s in margin:
            changes.append({"type": "dropped", "text": s, "reason": "Header/footer content skipped."})
            lines.append({"text": s, "status": "dropped"})
            continue
        cleaned = PUA_RE.sub("", EMOJI_RE.sub("", s)).strip() if p["drops_icons"] else s
        if cleaned != s:
            changes.append({"type": "icon", "text": s, "reason": "Icon glyphs stripped."})
        key = re.sub(r"\s+", " ", re.sub(r"[^a-z&' ]", " ", cleaned.lower())).strip()
        if p["strict_headers"] and key in NONSTANDARD_SECTIONS and len(cleaned.split()) <= 4:
            changes.append({"type": "header", "text": cleaned, "reason": f"'{cleaned}' isn't a recognised section; content may be misfiled."})
            lines.append({"text": cleaned, "status": "unrecognized"})
            continue
        lines.append({"text": cleaned, "status": "changed" if cleaned != s else "ok"})
    return {"lines": lines[:400], "changes": changes[:60]}


def read_resume(data: bytes, filename: str) -> Dict[str, Any]:
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:4] == b"%PDF":
        read = _read_pdf(data)
    elif name.endswith(".docx"):
        read = _read_docx(data)
    else:
        text = data.decode("utf-8", errors="ignore")
        read = {"text": text, "naive_text": text, "file_type": "text",
                "layout": {"pages": None, "column_pages": 0, "tables": 0, "images": 0, "icons": 0, "margin_contact": [], "fonts": []}}
    read["text"] = re.sub(r"[ \t]+\n", "\n", read["text"]).strip()
    return read


def read_text(text: str) -> Dict[str, Any]:
    return {"text": text.strip(), "naive_text": text.strip(), "file_type": "text",
            "layout": {"pages": None, "column_pages": 0, "tables": 0, "images": 0, "icons": 0, "margin_contact": [], "fonts": []}}
