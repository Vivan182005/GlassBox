"""
Tailored, ATS-safe resume for one job.

The AI restructures the candidate's own resume around the job's keywords;
a fact guard in code then reverts any bullet that introduces a number the
original resume doesn't contain and drops employers that aren't in it.
The result is re-scored with the same ATS check, so the before/after score
is measured, not promised. Output is single column, standard headings,
plain fonts and full URLs: everything the X-ray penalises is designed out.
"""
import io
import re
from typing import Any, Dict, List, Optional

from core.llm import llm_json
from matching.ats_check import run_check
from matching.skills_lexicon import implied_by, term_in_text
from parsing.resume_reader import audit, read_text

SECTION_ORDER = ["summary", "skills", "experience", "projects", "education", "certifications", "achievements"]
HEADINGS = {"summary": "SUMMARY", "skills": "SKILLS", "experience": "EXPERIENCE", "projects": "PROJECTS",
            "education": "EDUCATION", "certifications": "CERTIFICATIONS", "achievements": "ACHIEVEMENTS"}

_PROMPT = (
    "You rewrite a candidate's resume for ONE job so it passes ATS keyword screening and reads well to a recruiter. "
    "Hard rules:\n"
    "1. Use ONLY facts in the original resume. Never add employers, titles, dates, degrees, numbers, tools or claims "
    "that aren't there. Rephrasing and reordering is fine; inventing is not.\n"
    "2. Weave in the job's keywords ONLY where the resume genuinely supports them, using the job's exact wording "
    "(e.g. if the resume says 'dashboards in Tableau' and the JD says 'data visualization', you may say "
    "'data visualization dashboards in Tableau'). List unsupported keywords in do_not_claim.\n"
    "3. Bullets: start with a strong verb, one line or two, keep every number from the original bullet.\n"
    "4. Order skills with the job's must-haves first. Order projects by relevance to the job.\n"
    "5. headline: the job's title wording if truthful for this candidate's level, else their real target role.\n"
    "Return ONLY JSON:\n"
    '{"headline": str, "summary": "2-3 sentences", '
    '"skills": [{"group": str, "items": [str]}], '
    '"experience": [{"title": str, "company": str, "location": str|null, "dates": str, '
    '"bullets": [{"text": str, "source": "the original bullet this came from, verbatim"}]}], '
    '"projects": [{"name": str, "tech": str|null, "link": str|null, "bullets": [{"text": str, "source": str}]}], '
    '"education": [{"degree": str, "school": str, "dates": str|null, "detail": str|null}], '
    '"certifications": [str], "achievements": [str], '
    '"do_not_claim": [str], '
    '"cover_letter": "150-220 words, first person, specific to this company and role, only resume facts, no greeting or sign-off"}'
)

_NUM = re.compile(r"\d+(?:[.,]\d+)?")


def _numbers(text: str) -> set:
    return {n.replace(",", "") for n in _NUM.findall(text or "")}


def _normalize(data: Dict[str, Any]) -> Dict[str, Any]:
    """LLMs sometimes return a list of strings where objects were asked for;
    coerce every section to the documented shape so rendering never breaks."""
    def as_list(v):
        return v if isinstance(v, list) else ([v] if v else [])

    skills = []
    for g in as_list(data.get("skills")):
        if isinstance(g, str):
            skills.append({"group": "Skills", "items": [x.strip() for x in g.split(",") if x.strip()]})
        elif isinstance(g, dict):
            skills.append({"group": g.get("group") or "Skills", "items": [str(i) for i in as_list(g.get("items")) if i]})
    data["skills"] = skills
    data["experience"] = [e for e in as_list(data.get("experience")) if isinstance(e, dict)]
    data["projects"] = [p if isinstance(p, dict) else {"name": str(p), "bullets": []} for p in as_list(data.get("projects"))]
    data["education"] = [e if isinstance(e, dict) else {"degree": str(e), "school": ""} for e in as_list(data.get("education"))]
    for key in ("certifications", "achievements", "do_not_claim"):
        data[key] = [str(x) for x in as_list(data.get(key)) if x and not isinstance(x, (dict, list))]
    for key in ("headline", "summary", "cover_letter"):
        if not isinstance(data.get(key), str):
            data[key] = ""
    return data


def _guard(data: Dict[str, Any], source: str) -> List[str]:
    """Reverts bullets with numbers the original resume doesn't have; drops
    employers that aren't in it. Returns notes on what was corrected."""
    src_nums = _numbers(source)
    src_lower = re.sub(r"\s+", " ", source.lower())
    notes = []

    def fix_bullets(bullets):
        out = []
        for b in bullets or []:
            if isinstance(b, str):
                b = {"text": b, "source": ""}
            text = (b.get("text") or "").strip()
            extra = _numbers(text) - src_nums
            if extra:
                orig = (b.get("source") or "").strip()
                if orig and not (_numbers(orig) - src_nums):
                    notes.append(f"Kept your original wording where the rewrite added {', '.join(sorted(extra))}.")
                    text = orig
                else:
                    notes.append(f"Dropped a bullet that claimed {', '.join(sorted(extra))}, which isn't on your resume.")
                    continue
            if text:
                out.append(text)
        return out

    # Skills: every item must be traceable to the original resume, and nothing
    # the model itself flagged as unsupported may slip back in.
    # The resume is the authority in both directions: a skill it shows is kept
    # even if the model listed it under do_not_claim, and do_not_claim is
    # cleaned of anything the resume actually has.
    def supported(item: str) -> bool:
        item = item.strip()
        return bool(item) and (term_in_text(item, source) > 0 or bool(implied_by(item, source)))

    data["do_not_claim"] = [t for t in data.get("do_not_claim") or [] if not supported(t)]

    for g in data.get("skills") or []:
        items = []
        for raw in g.get("items") or []:
            m = re.match(r"^(.*?)\s*\((.*)\)\s*$", raw or "")
            if m:  # "AI tools (LLMs, OpenAI API, Claude)": check each inner item
                inner = [x.strip() for x in m.group(2).split(",") if supported(x)]
                base_ok = supported(m.group(1))
                if inner:
                    items.append(f"{m.group(1)} ({', '.join(inner)})")
                elif base_ok:
                    items.append(m.group(1))
                dropped = [x.strip() for x in m.group(2).split(",") if x.strip() and x.strip() not in inner]
            else:
                dropped = [] if supported(raw or "") else [raw]
                if not dropped:
                    items.append(raw)
            for d in dropped:
                if d:
                    notes.append(f"Removed skill '{d}', which isn't on your resume.")
        g["items"] = items

    kept = []
    for e in data.get("experience") or []:
        company = (e.get("company") or "").strip()
        first = company.lower().split(",")[0].split("(")[0].strip()
        if first and first[:12] not in src_lower:
            notes.append(f"Removed '{company}', which isn't on your original resume.")
            continue
        e["bullets"] = fix_bullets(e.get("bullets"))
        kept.append(e)
    data["experience"] = kept
    for p in data.get("projects") or []:
        p["bullets"] = fix_bullets(p.get("bullets"))
    return notes


def to_text(data: Dict[str, Any], header: Dict[str, Any]) -> str:
    """Plain-text rendering; also what the ATS check scores."""
    lines = [header["name"]]
    if data.get("headline"):
        lines.append(data["headline"])
    lines.append(" | ".join(header["contact"]))
    for sec in SECTION_ORDER:
        body = _section_lines(data, sec)
        if body:
            lines += ["", HEADINGS[sec]] + body
    return "\n".join(lines)


def _section_lines(data: Dict[str, Any], sec: str) -> List[str]:
    out: List[str] = []
    if sec == "summary" and data.get("summary"):
        out.append(data["summary"])
    elif sec == "skills":
        for g in data.get("skills") or []:
            items = [i for i in g.get("items") or [] if i]
            if items:
                out.append(f"{g.get('group') or 'Skills'}: {', '.join(items)}")
    elif sec == "experience":
        for e in data.get("experience") or []:
            out.append(f"{e.get('title', '')}, {e.get('company', '')}" + (f", {e['location']}" if e.get("location") else "")
                       + (f" | {e['dates']}" if e.get("dates") else ""))
            out += [f"- {b}" for b in e.get("bullets") or []]
    elif sec == "projects":
        for p in data.get("projects") or []:
            out.append(p.get("name", "") + (f" | {p['tech']}" if p.get("tech") else "") + (f" | {p['link']}" if p.get("link") else ""))
            out += [f"- {b}" for b in p.get("bullets") or []]
    elif sec == "education":
        for ed in data.get("education") or []:
            out.append(", ".join(x for x in [ed.get("degree"), ed.get("school")] if x) + (f" | {ed['dates']}" if ed.get("dates") else ""))
            if ed.get("detail"):
                out.append(ed["detail"])
    elif sec in ("certifications", "achievements"):
        out += [f"- {c}" for c in data.get(sec) or [] if c]
    return out


def _link_projects(data: Dict[str, Any], link_lines: List[Dict[str, str]]) -> None:
    """Replace placeholder link text ("GitHub", "Link") with the real URL the
    original PDF hid behind it. Each link is matched through the line it sat
    on in the original (the project's title row); ATS can't follow links."""
    for p in data.get("projects") or []:
        link = (p.get("link") or "").strip()
        if re.match(r"^(https?://|www\.)", link) or re.search(r"\.(com|io|dev|app|in)/", link):
            continue
        words = {w for w in re.findall(r"[a-z0-9]+", (p.get("name") or "").lower()) if len(w) > 2}
        best, best_hits = None, 0
        for l in link_lines:
            if re.search(r"linkedin\.com|mailto:", l["uri"]):
                continue
            line = re.sub(r"[^a-z0-9]+", " ", l["line"].lower())
            hits = sum(1 for w in words if re.search(r"\b" + re.escape(w), line))
            if hits > best_hits:
                best, best_hits = l["uri"], hits
        p["link"] = re.sub(r"^https?://(www\.)?", "", best).rstrip("/") if best else None


def _header(profile: Dict[str, Any], contact: Dict[str, Any]) -> Dict[str, Any]:
    bits = []
    if profile.get("location"):
        bits.append(f"{profile['location']}, India")
    for k in ("phone", "email", "linkedin", "github"):
        if contact.get(k):
            bits.append(re.sub(r"^https?://(www\.)?", "", contact[k]).rstrip("/"))
    return {"name": profile.get("name") or "Your Name", "contact": bits}


def build(resume_text: str, profile: Dict[str, Any], audit_result: Dict[str, Any], jd_text: str, before: Dict[str, Any],
          ats_id: str, job_title: str, company: str) -> Optional[Dict[str, Any]]:
    must = [r["term"] for r in before["keywords"]["must_have"]]
    implied = [f"{r['term']} (shown by {', '.join(r['via'])})" for r in before["keywords"]["must_have"] + before["keywords"]["nice_to_have"]
               if r["status"] == "implied"]
    nice = [r["term"] for r in before["keywords"]["nice_to_have"]]
    data = llm_json(_PROMPT,
                    f"JOB: {job_title} at {company}\nMUST-HAVE KEYWORDS: {', '.join(must)}\nNICE-TO-HAVE: {', '.join(nice)}\n"
                    f"IMPLIED BY THE RESUME (name these explicitly, they are truthful): {', '.join(implied) or 'none'}\n"
                    f"CANDIDATE LEVEL: {profile.get('seniority')} ({profile.get('years_experience')} years incl. internships)\n\n"
                    f"JOB DESCRIPTION:\n{jd_text[:5000]}\n\nORIGINAL RESUME:\n{resume_text[:9000]}",
                    max_tokens=4500, timeout=90)
    if not data:
        return None
    notes = _guard(_normalize(data), resume_text)
    _link_projects(data, (audit_result.get("contact") or {}).get("link_lines") or [])
    header = _header(profile, audit_result.get("contact") or {})
    text = to_text(data, header)

    new_audit = audit(read_text(text))
    after = run_check(text, profile, new_audit, jd_text, ats_id, job_title, company)
    return {
        "resume": data, "header": header, "text": text, "fact_notes": notes,
        "do_not_claim": data.get("do_not_claim") or [],
        "cover_letter": data.get("cover_letter") or "",
        "before": {"score": before["score"], "verdict": before["verdict"], "coverage": before["coverage"]},
        "after": {"score": after["score"], "verdict": after["verdict"], "coverage": after["coverage"],
                  "keywords": after["keywords"], "fixes": after["fixes"]},
        "audit_issues": [i["title"] for i in new_audit["issues"]],
    }


# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------

_LATIN = str.maketrans({"–": "-", "—": "-", "‘": "'", "’": "'", "“": '"', "”": '"', "•": "-", "…": "...", "₹": "Rs.",
                        "→": "->", " ": " ", "​": ""})


def _latin(s: str) -> str:
    import unicodedata
    s = re.sub(r"[\u2010-\u2015\u2212]", "-", (s or "").translate(_LATIN)).replace("\u00ad", "")
    # Anything still outside Latin-1 (the PDF core fonts' range) is decomposed
    # to its base letter, or dropped, rather than printed as "?".
    out = []
    for ch in s:
        if ord(ch) < 256:
            out.append(ch)
        else:
            base = unicodedata.normalize("NFKD", ch).encode("latin-1", "ignore").decode("latin-1")
            out.append(base)
    return "".join(out)


def to_pdf(data: Dict[str, Any], header: Dict[str, Any]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_margins(16, 14, 16)
    pdf.set_auto_page_break(True, 14)
    pdf.add_page()
    w = pdf.w - pdf.l_margin - pdf.r_margin
    pdf.set_font("Helvetica", "B", 17)
    pdf.cell(w, 8, _latin(header["name"]), new_x="LMARGIN", new_y="NEXT", align="C")
    if data.get("headline"):
        pdf.set_font("Helvetica", "", 10.5)
        pdf.cell(w, 5.5, _latin(data["headline"]), new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.set_font("Helvetica", "", 9.5)
    pdf.multi_cell(w, 4.8, _latin(" | ".join(header["contact"])), align="C", new_x="LMARGIN", new_y="NEXT")

    def heading(t):
        pdf.ln(2.5)
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(w, 6, t, new_x="LMARGIN", new_y="NEXT")
        y = pdf.get_y()
        pdf.line(pdf.l_margin, y, pdf.l_margin + w, y)
        pdf.ln(1.2)

    def para(t, style="", size=10, h=4.9):
        pdf.set_font("Helvetica", style, size)
        pdf.multi_cell(w, h, _latin(t), align="L", new_x="LMARGIN", new_y="NEXT")

    def bullet(t):
        pdf.set_font("Helvetica", "", 10)
        pdf.set_x(pdf.l_margin + 3)
        pdf.multi_cell(w - 3, 4.9, _latin(f"- {t}"), align="L", new_x="LMARGIN", new_y="NEXT")

    for sec in SECTION_ORDER:
        if not _section_lines(data, sec):
            continue
        heading(HEADINGS[sec])
        if sec == "summary":
            para(data["summary"])
        elif sec == "skills":
            for g in data.get("skills") or []:
                if g.get("items"):
                    para(f"{g.get('group') or 'Skills'}: {', '.join(g['items'])}")
        elif sec == "experience":
            for e in data["experience"]:
                para(f"{e.get('title', '')}, {e.get('company', '')}" + (f", {e['location']}" if e.get("location") else "")
                     + (f"  |  {e['dates']}" if e.get("dates") else ""), "B")
                for b in e.get("bullets") or []:
                    bullet(b)
                pdf.ln(1)
        elif sec == "projects":
            for p in data["projects"]:
                para(p.get("name", "") + (f"  |  {p['tech']}" if p.get("tech") else ""), "B")
                if p.get("link"):
                    para(p["link"], size=9)
                for b in p.get("bullets") or []:
                    bullet(b)
                pdf.ln(1)
        elif sec == "education":
            for ed in data["education"]:
                para(", ".join(x for x in [ed.get("degree"), ed.get("school")] if x) + (f"  |  {ed['dates']}" if ed.get("dates") else ""), "B")
                if ed.get("detail"):
                    para(ed["detail"])
        else:
            for c in data.get(sec) or []:
                bullet(c)
    return bytes(pdf.output())


def to_docx(data: Dict[str, Any], header: Dict[str, Any]) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for s in d.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(1.8)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(header["name"])
    r.bold, r.font.size = True, Pt(17)
    for line in ([data["headline"]] if data.get("headline") else []) + [" | ".join(header["contact"])]:
        q = d.add_paragraph(line)
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for sec in SECTION_ORDER:
        body = _section_lines(data, sec)
        if not body:
            continue
        h = d.add_paragraph()
        hr = h.add_run(HEADINGS[sec])
        hr.bold, hr.font.size = True, Pt(11.5)
        for line in body:
            if line.startswith("- "):
                d.add_paragraph(line[2:], style="List Bullet")
            else:
                para = d.add_paragraph()
                run = para.add_run(line)
                run.bold = sec in ("experience", "projects", "education") and " | " in line
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
